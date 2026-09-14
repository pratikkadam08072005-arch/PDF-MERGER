import os
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document
from flask import Flask, jsonify, request, send_file
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
MAX_FILE_SIZE = 50 * 1024 * 1024
MAX_TOTAL_SIZE = 250 * 1024 * 1024
PDF_EXTENSIONS = {".pdf"}
WORD_EXTENSIONS = {".docx", ".doc"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}

app = Flask(__name__, static_folder=str(BASE_DIR), static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = MAX_TOTAL_SIZE


@app.get("/")
def index():
    return app.send_static_file("index.html")


@app.get("/sitemap.xml")
def sitemap():
    page_url = request.host_url.rstrip("/") + "/"
    xml = f'''<?xml version="1.0" encoding="UTF-8"?>
<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
  <url><loc>{page_url}</loc><changefreq>weekly</changefreq><priority>1.0</priority></url>
</urlset>'''
    return app.response_class(xml, mimetype="application/xml")


def uploaded_files():
    files = request.files.getlist("files")
    if not files:
        raise ValueError("Add at least one file to continue.")

    loaded = []
    total_size = 0
    for uploaded_file in files:
        filename = secure_filename(uploaded_file.filename or "unnamed")
        if not filename:
            raise ValueError("One uploaded file has no usable filename.")
        data = uploaded_file.read()
        if len(data) > MAX_FILE_SIZE:
            raise ValueError(f"{filename} is larger than 50 MB.")
        total_size += len(data)
        if total_size > MAX_TOTAL_SIZE:
            raise ValueError("The combined upload is larger than 250 MB.")
        loaded.append((filename, data))
    return loaded


def require_extensions(files, allowed, label):
    for filename, _data in files:
        if Path(filename).suffix.lower() not in allowed:
            raise ValueError(f"{filename} is not a supported {label} file.")


def pdf_response(data, filename):
    return send_file(BytesIO(data), mimetype="application/pdf", as_attachment=True, download_name=filename)


def docx_response(data, filename):
    return send_file(
        BytesIO(data),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=filename,
    )


def image_response(data, filename, mimetype):
    return send_file(BytesIO(data), mimetype=mimetype, as_attachment=True, download_name=filename)


@app.post("/merge")
def merge_pdfs():
    try:
        files = uploaded_files()
        if len(files) < 2:
            raise ValueError("Add at least two PDF files.")
        require_extensions(files, PDF_EXTENSIONS, "PDF")

        writer = PdfWriter()
        for _filename, data in files:
            reader = PdfReader(BytesIO(data))
            for page in reader.pages:
                writer.add_page(page)
        output = BytesIO()
        writer.write(output)
        return pdf_response(output.getvalue(), "paperfold-merged.pdf")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("PDF merge failed")
        return jsonify(error="One of the files could not be read as a standard PDF."), 400


@app.post("/pdf-to-word")
def pdf_to_word():
    try:
        files = uploaded_files()
        if len(files) != 1:
            raise ValueError("Choose one PDF file for PDF to Word conversion.")
        require_extensions(files, PDF_EXTENSIONS, "PDF")

        from pdf2docx import Converter

        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.pdf"
            target = Path(folder) / "paperfold-converted.docx"
            source.write_bytes(files[0][1])
            converter = Converter(str(source))
            try:
                converter.convert(str(target))
            finally:
                converter.close()
            return docx_response(target.read_bytes(), target.name)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("PDF to Word conversion failed")
        return jsonify(error="This PDF could not be converted to Word."), 400


@app.post("/word-to-pdf")
def word_to_pdf():
    try:
        files = uploaded_files()
        if len(files) != 1:
            raise ValueError("Choose one Word file for Word to PDF conversion.")
        require_extensions(files, WORD_EXTENSIONS, "Word")

        office = shutil.which("soffice") or shutil.which("libreoffice")
        if not office:
            return jsonify(error="Word to PDF needs LibreOffice installed on the host."), 503

        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / secure_filename(files[0][0])
            source.write_bytes(files[0][1])
            result = subprocess.run(
                [office, "--headless", "--convert-to", "pdf", "--outdir", folder, str(source)],
                capture_output=True,
                text=True,
                timeout=90,
                check=False,
            )
            output = source.with_suffix(".pdf")
            if result.returncode != 0 or not output.exists():
                raise RuntimeError(result.stderr.strip() or "LibreOffice could not create a PDF.")
            return pdf_response(output.read_bytes(), "paperfold-converted.pdf")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except subprocess.TimeoutExpired:
        return jsonify(error="The Word document took too long to convert."), 504
    except Exception:
        app.logger.exception("Word to PDF conversion failed")
        return jsonify(error="This Word document could not be converted to PDF."), 400


@app.post("/image-to-pdf")
def image_to_pdf():
    try:
        files = uploaded_files()
        require_extensions(files, IMAGE_EXTENSIONS, "image")
        images = []
        for _filename, data in files:
            with Image.open(BytesIO(data)) as image:
                images.append(image.convert("RGB"))
        output = BytesIO()
        images[0].save(output, format="PDF", save_all=True, append_images=images[1:])
        for image in images:
            image.close()
        return pdf_response(output.getvalue(), "paperfold-images.pdf")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("Image to PDF conversion failed")
        return jsonify(error="One of these images could not be converted to PDF."), 400


@app.post("/image-to-word")
def image_to_word():
    try:
        files = uploaded_files()
        require_extensions(files, IMAGE_EXTENSIONS, "image")
        document = Document()
        section = document.sections[0]
        max_width = section.page_width - section.left_margin - section.right_margin
        for index, (filename, data) in enumerate(files):
            with Image.open(BytesIO(data)) as image:
                image.verify()
            with tempfile.NamedTemporaryFile(suffix=Path(filename).suffix, delete=False) as temporary:
                temporary.write(data)
                temporary_path = temporary.name
            try:
                document.add_picture(temporary_path, width=max_width)
            finally:
                Path(temporary_path).unlink(missing_ok=True)
            if index < len(files) - 1:
                document.add_page_break()
        output = BytesIO()
        document.save(output)
        return docx_response(output.getvalue(), "paperfold-images.docx")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("Image to Word conversion failed")
        return jsonify(error="One of these images could not be added to Word."), 400


@app.post("/pdf-editor")
def edit_pdf():
    try:
        files = uploaded_files()
        if len(files) != 1:
            raise ValueError("Choose one PDF file to edit.")
        require_extensions(files, PDF_EXTENSIONS, "PDF")
        reader = PdfReader(BytesIO(files[0][1]))
        remove_pages = {int(value) - 1 for value in request.form.get("remove_pages", "").split(",") if value.strip().isdigit()}
        rotation = int(request.form.get("rotation", "0")) % 360
        overlay_text = request.form.get("text", "").strip()
        overlay_page = max(1, int(request.form.get("text_page", "1"))) - 1
        text_x = max(0, int(request.form.get("text_x", "72")))
        text_y = max(0, int(request.form.get("text_y", "72")))
        text_size = max(6, min(72, int(request.form.get("text_size", "16"))))
        writer = PdfWriter()
        for index, page in enumerate(reader.pages):
            if index in remove_pages:
                continue
            if rotation:
                page.rotate(rotation)
            if overlay_text and index == overlay_page:
                overlay = BytesIO()
                page_width = float(page.mediabox.width)
                page_height = float(page.mediabox.height)
                overlay_canvas = canvas.Canvas(overlay, pagesize=(page_width, page_height))
                overlay_canvas.setFont("Helvetica", text_size)
                overlay_canvas.drawString(text_x, min(text_y, page_height - text_size), overlay_text)
                overlay_canvas.save()
                page.merge_page(PdfReader(BytesIO(overlay.getvalue())).pages[0])
            writer.add_page(page)
        if len(writer.pages) == 0:
            raise ValueError("Keep at least one page in the edited PDF.")
        output = BytesIO()
        writer.write(output)
        return pdf_response(output.getvalue(), "paperfold-edited.pdf")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("PDF editing failed")
        return jsonify(error="This PDF could not be edited."), 400


@app.post("/word-editor")
def edit_word():
    try:
        files = uploaded_files()
        if len(files) != 1 or Path(files[0][0]).suffix.lower() != ".docx":
            raise ValueError("Choose one DOCX file to edit.")
        document = Document(BytesIO(files[0][1]))
        edited_text = request.form.get("text", "")
        if edited_text:
            for paragraph in document.paragraphs:
                paragraph.text = ""
            for line in edited_text.splitlines():
                document.add_paragraph(line)
        output = BytesIO()
        document.save(output)
        return docx_response(output.getvalue(), "paperfold-edited.docx")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("Word editing failed")
        return jsonify(error="This DOCX file could not be edited."), 400


@app.post("/image-editor")
def edit_image():
    try:
        files = uploaded_files()
        if len(files) != 1:
            raise ValueError("Choose one image to edit.")
        require_extensions(files, IMAGE_EXTENSIONS, "image")
        rotation = int(request.form.get("rotation", "0")) % 360
        brightness = max(-100, min(100, int(request.form.get("brightness", "0"))))
        contrast = max(-100, min(100, int(request.form.get("contrast", "0"))))
        overlay_text = request.form.get("text", "").strip()
        text_x = max(0, int(request.form.get("text_x", "24")))
        text_y = max(0, int(request.form.get("text_y", "24")))
        text_size = max(8, min(160, int(request.form.get("text_size", "32"))))
        from PIL import ImageEnhance
        from PIL import ImageDraw, ImageFont

        with Image.open(BytesIO(files[0][1])) as source:
            image = source.convert("RGB")
            if rotation:
                image = image.rotate(rotation, expand=True)
            if brightness:
                image = ImageEnhance.Brightness(image).enhance(1 + brightness / 100)
            if contrast:
                image = ImageEnhance.Contrast(image).enhance(1 + contrast / 100)
            if overlay_text:
                drawer = ImageDraw.Draw(image)
                try:
                    font = ImageFont.truetype("arial.ttf", text_size)
                except OSError:
                    font = ImageFont.load_default()
                drawer.text((text_x, text_y), overlay_text, fill=(32, 33, 30), font=font)
            output = BytesIO()
            image.save(output, format="PNG")
        return image_response(output.getvalue(), "paperfold-edited.png", "image/png")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("Image editing failed")
        return jsonify(error="This image could not be edited."), 400


@app.errorhandler(413)
def request_too_large(_error):
    return jsonify(error="The combined upload is larger than 250 MB."), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=True)
